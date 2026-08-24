# -*- coding: utf-8 -*-
"""
将 Markdown 章节合并生成一份套用 Word 模板样式的 .docx。

用法:
    python build_docx.py --template <模板.docx> --input <md目录或文件...> \
                         --output <输出.docx> [--title "文档标题"]

特性:
- 标题 #/##/###/#### -> Heading 1~4（继承模板内置样式）
- 表格 -> 带边框 Word 表格
- ```mermaid 代码块 -> 联网渲染为 PNG 图片插入（Kroki 优先，mermaid.ink 兜底）
- 其他 ``` 代码块 -> 等宽字体代码块
- 图题（图X-Y ...）-> 居中小字
- **加粗** / `行内代码` -> 对应 run 格式

依赖: python-docx（pip install python-docx）。Mermaid 渲染需联网。
"""
import sys
import io
import re
import os
import glob
import base64
import argparse
import tempfile
import urllib.request
from copy import deepcopy

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')

from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

CAPTION_RE = re.compile(r'^(图|表)\s*\d+[-－]\d+')  # 图题（图X-Y）与表题（表X-Y）统一居中小字
REF_RE = re.compile(r'^\[\d+\]')
INLINE_RE = re.compile(r'(\*\*.+?\*\*|`[^`]+`|\[\d+\])')

# 字体规格（依据已验收的初期报告标准）
EN_FONT = 'Times New Roman'
CN_BODY = '宋体'
CN_HEAD = '黑体'
BODY_SIZE = 12.0     # 正文 小四
REF_SIZE = 10.5      # 参考文献 五号
TABLE_SIZE = 10.5    # 表格 五号
FIG_SIZE = 9.0       # 图题
TITLE_SIZE = 18.0    # 文档标题
HEADING_SIZE = {1: 15.0, 2: 14.0, 3: 14.0, 4: 14.0}  # 一级小三，其余四号

STYLE_BY_ID = {}
H1_SEEN = [0]  # 一级标题（章）计数：从第二个起段前分页，第一个紧跟文档标题
FIRST_H1_NEEDS_PAGE_BREAK = [False]
HEADER_PROTOTYPES = {}


def header_prototype(doc, attr):
    """Find a non-empty template header for a header variant."""
    for section in doc.sections:
        header = getattr(section, attr)
        if header.paragraphs and header.paragraphs[0].text.strip():
            return header.paragraphs[0]
    return None


def set_header_text(header, title, prototype=None):
    """Write a chapter title while retaining the template header's paragraph rule."""
    header.is_linked_to_previous = False
    p = header.paragraphs[0] if header.paragraphs else header.add_paragraph()
    ppr = deepcopy(prototype._p.pPr) if prototype is not None and prototype._p.pPr is not None else None
    rpr = (deepcopy(prototype.runs[0]._r.rPr)
           if prototype is not None and prototype.runs and prototype.runs[0]._r.rPr is not None else None)
    if ppr is not None:
        old = p._p.pPr
        if old is not None:
            p._p.remove(old)
        p._p.insert(0, ppr)
    p.clear()
    run = p.add_run(title)
    if rpr is not None:
        run._r.get_or_add_rPr().append(rpr)
    else:
        set_font(run, CN_BODY, 10.5)


def set_section_header(section, title):
    """Keep default, even and first-page headers from falling back to prior sections."""
    set_header_text(section.header, title, HEADER_PROTOTYPES.get('header'))
    set_header_text(section.even_page_header, title, HEADER_PROTOTYPES.get('even_page_header'))
    set_header_text(section.first_page_header, title, HEADER_PROTOTYPES.get('first_page_header'))


def start_chapter_section(doc, title):
    """每个一级标题从新页开始，并与终期范例一样使用本章页眉。"""
    if H1_SEEN[0] > 1:
        section = doc.add_section(WD_SECTION.NEW_PAGE)
    else:
        section = doc.sections[-1]
    set_section_header(section, title)


def get_style(style_id):
    return STYLE_BY_ID.get(style_id)


def get_heading_style(level):
    """兼容 Word 模板中 Heading 样式的不同内部 ID（如 `1` 或 `Heading1`）。"""
    expected_name = f'Heading {level}'
    for style in STYLE_BY_ID.values():
        if style.name == expected_name:
            return style
    return get_style(f'Heading{level}')


def get_heading_style_id(level):
    style = get_heading_style(level)
    return style.style_id if style is not None else f'Heading{level}'


def is_toc_style_id(style_id):
    """兼容 `TOC1`、`11` 等 Word 自动目录样式 ID。"""
    if not style_id:
        return False
    style = get_style(style_id)
    return style_id.upper().startswith('TOC') or (style is not None and style.name.lower().startswith('toc'))


def clear_body(doc):
    body = doc.element.body
    for child in list(body):
        if child.tag == qn('w:sectPr'):
            continue
        body.remove(child)


def trim_to_content_area(doc):
    """保留模板封面、目录及目录后的分页/分节符，删除模板的"说明"提示段与正文占位段，
    使生成的正文写入目录之后的内容区。无法识别模板结构时返回 False（交由调用方回退）。"""
    body = doc.element.body
    P, PPR, PSTYLE, VAL = qn('w:p'), qn('w:pPr'), qn('w:pStyle'), qn('w:val')
    SECTPR, PGBB, BR, TYPE = qn('w:sectPr'), qn('w:pageBreakBefore'), qn('w:br'), qn('w:type')

    def pstyle(el):
        if el.tag != P:
            return None
        s = el.find(PPR + '/' + PSTYLE)
        return s.get(VAL) if s is not None else ''

    children = list(body.iterchildren())
    doc_sectpr = body.find(SECTPR)

    h1_style_id = get_heading_style_id(1)
    first_h1 = next((el for el in children if pstyle(el) == h1_style_id), None)
    last_toc = None
    for el in children:
        st = pstyle(el)
        if is_toc_style_id(st):
            last_toc = el
    if first_h1 is None or last_toc is None:
        return False

    idx_h1 = children.index(first_h1)
    idx_toc = children.index(last_toc)
    if idx_toc >= idx_h1:
        return False

    # 目录与正文之间带分页/分节的分隔段（保留以维持目录后的分页与分节属性）
    sep_idx = idx_h1
    for i in range(idx_toc + 1, idx_h1):
        el = children[i]
        if el.tag != P:
            continue
        pPr = el.find(PPR)
        has_sep = pPr is not None and (pPr.find(PGBB) is not None or pPr.find(SECTPR) is not None)
        if not has_sep and el.find('.//' + BR + '[@' + TYPE + '="page"]') is not None:
            has_sep = True
        if has_sep:
            sep_idx = i

    # 删除目录之后、分隔段之前的"说明"提示段
    for el in children[idx_toc + 1: sep_idx]:
        body.remove(el)
    # 删除正文占位段（第一个一级标题起到末尾），保留文档级 sectPr
    for el in children[idx_h1:]:
        if el is doc_sectpr:
            continue
        body.remove(el)

    # 未找到分隔段则正文需自带首页分页：让第一个一级标题也段前分页
    return True if sep_idx != idx_h1 else 'no_sep'


def set_font(run, cn=CN_BODY, size=BODY_SIZE, bold=False, superscript=False, en=EN_FONT):
    """显式设置中英文字体、字号，确保与标准报告一致。"""
    run.bold = bold
    if size is not None:
        run.font.size = Pt(size)
    if superscript:
        run.font.superscript = True
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn('w:rFonts'))
    if rfonts is None:
        rfonts = OxmlElement('w:rFonts')
        rpr.append(rfonts)
    rfonts.set(qn('w:ascii'), en)
    rfonts.set(qn('w:hAnsi'), en)
    rfonts.set(qn('w:eastAsia'), cn)


def set_code_font(run):
    run.font.size = Pt(9)
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn('w:rFonts'))
    if rfonts is None:
        rfonts = OxmlElement('w:rFonts')
        rpr.append(rfonts)
    rfonts.set(qn('w:ascii'), 'Consolas')
    rfonts.set(qn('w:hAnsi'), 'Consolas')


def add_inline(paragraph, text, size=BODY_SIZE, cn=CN_BODY, sup_refs=True):
    """处理 **加粗**、`行内代码`、[n]上标引用，其余为普通 run。"""
    for part in INLINE_RE.split(text):
        if not part:
            continue
        if part.startswith('**') and part.endswith('**'):
            set_font(paragraph.add_run(part[2:-2]), cn, size, bold=True)
        elif part.startswith('`') and part.endswith('`'):
            set_code_font(paragraph.add_run(part[1:-1]))
        elif sup_refs and re.fullmatch(r'\[\d+\]', part):
            set_font(paragraph.add_run(part), cn, size, superscript=True)
        else:
            set_font(paragraph.add_run(part), cn, size)


def png_size(data):
    if data[:8] == b'\x89PNG\r\n\x1a\n':
        return int.from_bytes(data[16:20], 'big'), int.from_bytes(data[20:24], 'big')
    return None, None


def _kroki(code):
    req = urllib.request.Request(
        'https://kroki.io/mermaid/png', data=code.encode('utf-8'),
        headers={'Content-Type': 'text/plain', 'User-Agent': 'Mozilla/5.0'})
    with urllib.request.urlopen(req, timeout=45) as r:
        return r.read()


def _mermaid_ink(code):
    b64 = base64.urlsafe_b64encode(code.encode('utf-8')).decode('ascii')
    url = 'https://mermaid.ink/img/' + b64 + '?type=png&bgColor=ffffff'
    req = urllib.request.Request(url, headers={'User-Agent': 'Mozilla/5.0'})
    with urllib.request.urlopen(req, timeout=45) as r:
        return r.read()


def render_mermaid(code):
    # 优先 Kroki（POST 无 URL 长度限制），失败再退回 mermaid.ink
    try:
        return _kroki(code)
    except Exception:
        return _mermaid_ink(code)


def add_image(doc, png_bytes, max_cm=15.0):
    fd, tmp = tempfile.mkstemp(suffix='.png')
    os.close(fd)
    try:
        with open(tmp, 'wb') as f:
            f.write(png_bytes)
        w, _ = png_size(png_bytes)
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run()
        if w:
            cm_w = w / 96 * 2.54
            shape = run.add_picture(tmp, width=Cm(min(cm_w, max_cm)))
        else:
            shape = run.add_picture(tmp, width=Cm(max_cm))
        # A caption is added by the Markdown renderer. The image still needs a
        # non-empty description for assistive technologies and document audits.
        shape._inline.docPr.set('descr', '报告工程图')
        shape._inline.docPr.set('title', '报告工程图')
    finally:
        os.remove(tmp)


def set_code_box(paragraph):
    """给段落加四周边框（白底无填充），形成代码框。"""
    pPr = paragraph._p.get_or_add_pPr()
    pbdr = OxmlElement('w:pBdr')
    for edge in ('top', 'left', 'bottom', 'right'):
        el = OxmlElement(f'w:{edge}')
        el.set(qn('w:val'), 'single')
        el.set(qn('w:sz'), '4')
        el.set(qn('w:space'), '6')
        el.set(qn('w:color'), '999999')
        pbdr.append(el)
    pPr.append(pbdr)


def add_code_block(doc, lines):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.5)
    set_code_box(p)
    for i, ln in enumerate(lines):
        if i > 0:
            p.add_run().add_break()
        set_code_font(p.add_run(ln))


def set_table_borders(table):
    tblPr = table._tbl.tblPr
    borders = OxmlElement('w:tblBorders')
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        el = OxmlElement(f'w:{edge}')
        el.set(qn('w:val'), 'single')
        el.set(qn('w:sz'), '4')
        el.set(qn('w:space'), '0')
        el.set(qn('w:color'), '000000')
        borders.append(el)
    tblPr.append(borders)


def add_table(doc, rows):
    ncol = len(rows[0])
    table = doc.add_table(rows=0, cols=ncol)
    set_table_borders(table)
    for ri, row in enumerate(rows):
        cells = table.add_row().cells
        for ci in range(ncol):
            txt = (row[ci] if ci < len(row) else '').replace('**', '').replace('`', '')
            cells[ci].text = ''
            run = cells[ci].paragraphs[0].add_run(txt)
            set_font(run, CN_BODY, TABLE_SIZE, bold=(ri == 0))
    tr_pr = table.rows[0]._tr.get_or_add_trPr()
    header = OxmlElement('w:tblHeader')
    header.set(qn('w:val'), 'true')
    tr_pr.append(header)
    return table


def request_word_field_update(doc):
    """Ask Word/WPS to refresh TOC and page fields when the file is opened."""
    settings = doc.settings.element
    node = settings.find(qn('w:updateFields'))
    if node is None:
        node = OxmlElement('w:updateFields')
        settings.append(node)
    node.set(qn('w:val'), 'true')


def split_table_row(line):
    s = line.strip().strip('|')
    return [c.strip() for c in s.split('|')]


def is_separator_row(line):
    return bool(re.match(r'^\s*\|?\s*:?-{2,}', line)) and set(line.strip()) <= set('|-: ')


def collect_md(inputs):
    """inputs 可为目录（取其中 *.md 排序）或文件路径，按给定顺序合并。"""
    files = []
    for p in inputs:
        if os.path.isdir(p):
            files.extend(sorted(glob.glob(os.path.join(p, '*.md'))))
        else:
            files.append(p)
    return files


def render_md(doc, lines, fig_counter):
    i, n = 0, len(lines)
    while i < n:
        line = lines[i]
        stripped = line.strip()

        if stripped.startswith('```'):
            lang = stripped[3:].strip().lower()
            buf = []
            i += 1
            while i < n and not lines[i].strip().startswith('```'):
                buf.append(lines[i])
                i += 1
            i += 1
            if lang == 'mermaid':
                fig_counter[0] += 1
                idx = fig_counter[0]
                try:
                    add_image(doc, render_mermaid('\n'.join(buf)))
                    print(f'  [图片] 第{idx}张 mermaid 渲染成功')
                except Exception as e:
                    print(f'  [警告] 第{idx}张 mermaid 渲染失败: {e}，改插文本')
                    add_code_block(doc, buf)
            else:
                add_code_block(doc, buf)
            continue

        if not stripped:
            i += 1
            continue

        m = re.match(r'^(#{1,4})\s+(.*)$', stripped)
        if m:
            level = len(m.group(1))
            if level == 1:
                H1_SEEN[0] += 1
                start_chapter_section(doc, m.group(2))
                print(f'  [章] 第{H1_SEEN[0]}个一级标题 "{m.group(2)[:16]}" 新建分节={H1_SEEN[0] > 1}')
            p = doc.add_paragraph()
            hs = get_heading_style(level)
            if hs is not None:
                p.style = hs
            if level == 1 and H1_SEEN[0] == 1 and FIRST_H1_NEEDS_PAGE_BREAK[0]:
                pPr = p._p.get_or_add_pPr()
                pPr.insert(0, OxmlElement('w:pageBreakBefore'))
            set_font(p.add_run(m.group(2)), CN_HEAD, HEADING_SIZE.get(level, 14.0))
            i += 1
            continue

        if stripped.startswith('|'):
            tbl = [split_table_row(line)]
            i += 1
            while i < n and lines[i].strip().startswith('|'):
                if not is_separator_row(lines[i]):
                    tbl.append(split_table_row(lines[i]))
                i += 1
            add_table(doc, tbl)
            continue

        if CAPTION_RE.match(stripped):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            set_font(p.add_run(stripped), CN_BODY, FIG_SIZE)
            i += 1
            continue

        # 参考文献条目：五号字，开头 [n] 不作上标
        if REF_RE.match(stripped):
            p = doc.add_paragraph()
            add_inline(p, stripped, size=REF_SIZE, sup_refs=False)
            i += 1
            continue

        if stripped.startswith('- '):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.74)
            set_font(p.add_run('• '), CN_BODY, BODY_SIZE)
            add_inline(p, stripped[2:])
            i += 1
            continue

        p = doc.add_paragraph()
        add_inline(p, stripped)
        i += 1


def build(template, inputs, output, title=None, rebuild=False):
    doc = Document(template)
    STYLE_BY_ID.clear()
    H1_SEEN[0] = 0
    FIRST_H1_NEEDS_PAGE_BREAK[0] = False
    HEADER_PROTOTYPES.clear()
    for attr in ('header', 'even_page_header', 'first_page_header'):
        prototype = header_prototype(doc, attr)
        if prototype is not None:
            HEADER_PROTOTYPES[attr] = prototype
    for s in doc.styles:
        STYLE_BY_ID[s.style_id] = s

    kept = False
    if not rebuild:
        kept = trim_to_content_area(doc)
        if kept == 'no_sep':
            FIRST_H1_NEEDS_PAGE_BREAK[0] = True
            kept = True
        if kept:
            print('已保留模板封面与目录，正文写入内容区')
        else:
            print('未识别到模板封面/目录结构，回退为整体重建')

    if not kept:
        clear_body(doc)
        if title:
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            set_font(p.add_run(title), CN_HEAD, TITLE_SIZE, bold=True)

    files = collect_md(inputs)
    if not files:
        raise SystemExit('未找到任何输入 md 文件')
    fig_counter = [0]
    for path in files:
        print(f'处理: {path}')
        with open(path, encoding='utf-8') as f:
            render_md(doc, f.read().split('\n'), fig_counter)

    request_word_field_update(doc)
    doc.save(output)
    print(f'\n已生成: {output}（共 {fig_counter[0]} 张图）。请在 Word/WPS 更新域并按最终分页核验目录页码。')


def main():
    parser = argparse.ArgumentParser(description='Markdown 章节 + Word 模板 -> 排版好的 .docx')
    parser.add_argument('--template', '-t', required=True, help='Word 模板 .docx 路径（提供标题/正文样式）')
    parser.add_argument('--input', '-i', nargs='+', required=True, help='章节 md 目录或文件列表（按顺序合并）')
    parser.add_argument('--output', '-o', required=True, help='输出 .docx 路径')
    parser.add_argument('--title', default=None, help='可选：文档首行居中标题（仅 --rebuild 整体重建时生效）')
    parser.add_argument('--rebuild', action='store_true',
                        help='整体重建：清空模板正文（含封面/目录）后重写；默认保留模板封面与目录，仅写入正文内容区')
    args = parser.parse_args()
    build(args.template, args.input, args.output, args.title, args.rebuild)


if __name__ == '__main__':
    main()
